from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
import re
from typing import Any


PROTOCOL_VERSION = 1
SUPPORTED_EXTENSIONS = {
    ".csv": "csv",
    ".docx": "docx",
    ".htm": "html",
    ".html": "html",
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".txt": "txt",
    ".xlsx": "xlsx",
}


class ConversionError(Exception):
    def __init__(self, code: str, message: str, *, retryable: bool = False):
        super().__init__(message)
        self.code = code
        self.message = message
        self.retryable = retryable


@dataclass(frozen=True)
class ConversionResult:
    markdown: str
    warnings: list[str]


@dataclass(frozen=True)
class WorkbookSheetInfo:
    name: str
    index: int
    visibility: str
    is_empty: bool
    row_count: int
    column_count: int
    preview: list[list[str | None]]


@dataclass(frozen=True)
class WorkbookInspection:
    sheets: list[WorkbookSheetInfo]


def extract_markdown(result: Any) -> str:
    text_content = getattr(result, "text_content", None)
    if isinstance(text_content, str):
        return text_content

    markdown = getattr(result, "markdown", None)
    if isinstance(markdown, str):
        return markdown

    if isinstance(result, str):
        return result

    raise ConversionError(
        "conversion_failed",
        "MarkItDown returned an unsupported result type",
    )


def validate_input(raw_path: str, requested_type: str | None = None) -> Path:
    input_path = Path(raw_path).expanduser()
    if not input_path.is_absolute():
        raise ConversionError("invalid_input", "inputPath must be absolute")
    if not input_path.exists():
        raise ConversionError("unreadable_file", f"file not found: {input_path}")
    if not input_path.is_file():
        raise ConversionError("invalid_input", f"not a file: {input_path}")

    detected_type = SUPPORTED_EXTENSIONS.get(input_path.suffix.lower())
    if detected_type is None:
        raise ConversionError(
            "unsupported_type",
            f"unsupported file type: {input_path.suffix or '(none)'}",
        )
    if requested_type and requested_type != detected_type:
        raise ConversionError(
            "invalid_input",
            f"inputType {requested_type!r} does not match {detected_type!r}",
        )
    return input_path.resolve()


class Converter:
    def __init__(self) -> None:
        self._converter: Any | None = None

    def _get_converter(self) -> Any:
        if self._converter is None:
            from markitdown import MarkItDown

            self._converter = MarkItDown()
        return self._converter

    def inspect_workbook(
        self, raw_path: str, requested_type: str | None = None
    ) -> WorkbookInspection:
        input_path = validate_input(raw_path, requested_type)
        if input_path.suffix.lower() != ".xlsx":
            raise ConversionError(
                "invalid_input", "workbook inspection is only available for xlsx"
            )

        try:
            from openpyxl import load_workbook

            workbook = load_workbook(input_path, read_only=True, data_only=True)
            try:
                sheets = [
                    self._inspect_sheet(worksheet, index)
                    for index, worksheet in enumerate(workbook.worksheets)
                ]
            finally:
                workbook.close()
            return WorkbookInspection(sheets=sheets)
        except ConversionError:
            raise
        except PermissionError as error:
            raise ConversionError("unreadable_file", str(error)) from error
        except Exception as error:
            raise self._map_conversion_error(error) from error

    def convert(
        self,
        raw_path: str,
        requested_type: str | None = None,
        sheet_names: list[str] | None = None,
    ) -> ConversionResult:
        input_path = validate_input(raw_path, requested_type)
        try:
            if sheet_names is not None:
                if input_path.suffix.lower() != ".xlsx":
                    raise ConversionError(
                        "invalid_input", "sheetNames is only available for xlsx"
                    )
                markdown = self._convert_xlsx(input_path, sheet_names)
            else:
                result = self._get_converter().convert(str(input_path))
                markdown = extract_markdown(result)
        except ConversionError:
            raise
        except PermissionError as error:
            raise ConversionError("unreadable_file", str(error)) from error
        except Exception as error:
            raise self._map_conversion_error(error) from error

        warnings: list[str] = []
        if input_path.suffix.lower() == ".pdf":
            without_comments = re.sub(r"<!--.*?-->", "", markdown, flags=re.DOTALL)
            visible_text_length = sum(
                1 for character in without_comments if character.isalnum()
            )
            if visible_text_length == 0:
                raise ConversionError(
                    "no_extractable_text",
                    "PDF contains no extractable text; OCR is not available",
                )
            if visible_text_length < 50:
                warnings.append("possible_scanned_pdf")

        return ConversionResult(markdown=markdown, warnings=warnings)

    @staticmethod
    def _inspect_sheet(worksheet: Any, index: int) -> WorkbookSheetInfo:
        row_count = worksheet.max_row or 1
        column_count = worksheet.max_column or 1
        if row_count == 1 and column_count == 1:
            reset_dimensions = getattr(worksheet, "reset_dimensions", None)
            if callable(reset_dimensions):
                reset_dimensions()
                return Converter._scan_sheet(worksheet, index)
        is_empty = (
            row_count == 1
            and column_count == 1
            and worksheet.cell(row=1, column=1).value is None
        )
        preview: list[list[str | None]] = []
        if not is_empty:
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(row_count, 20),
                min_col=1,
                max_col=min(column_count, 12),
                values_only=True,
            ):
                preview.append([Converter._preview_value(value) for value in row])
        return WorkbookSheetInfo(
            name=worksheet.title,
            index=index,
            visibility=worksheet.sheet_state,
            is_empty=is_empty,
            row_count=0 if is_empty else row_count,
            column_count=0 if is_empty else column_count,
            preview=preview,
        )

    @staticmethod
    def _scan_sheet(worksheet: Any, index: int) -> WorkbookSheetInfo:
        row_count = 0
        column_count = 0
        preview_rows: list[list[Any]] = []
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = list(row)
            populated_columns = [
                column_index
                for column_index, value in enumerate(values, start=1)
                if value is not None
            ]
            if populated_columns:
                row_count = row_index
                column_count = max(column_count, populated_columns[-1])
            if row_index <= 20:
                preview_rows.append(values[:12])

        is_empty = row_count == 0
        preview_width = min(column_count, 12)
        preview = (
            []
            if is_empty
            else [
                [
                    Converter._preview_value(
                        row[column_index] if column_index < len(row) else None
                    )
                    for column_index in range(preview_width)
                ]
                for row in preview_rows[: min(row_count, 20)]
            ]
        )
        return WorkbookSheetInfo(
            name=worksheet.title,
            index=index,
            visibility=worksheet.sheet_state,
            is_empty=is_empty,
            row_count=row_count,
            column_count=column_count,
            preview=preview,
        )

    @staticmethod
    def _preview_value(value: Any) -> str | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            text = value.isoformat(sep=" ", timespec="seconds")
        elif isinstance(value, (date, time)):
            text = value.isoformat()
        else:
            text = str(value)
        return text if len(text) <= 200 else f"{text[:197]}..."

    @staticmethod
    def _convert_xlsx(input_path: Path, sheet_names: list[str]) -> str:
        if not sheet_names or len(sheet_names) != len(set(sheet_names)):
            raise ConversionError(
                "sheet_selection_invalid",
                "at least one unique worksheet must be selected",
            )

        from openpyxl import load_workbook
        import pandas as pd
        from markitdown.converters import HtmlConverter

        workbook = load_workbook(input_path, read_only=True, data_only=True)
        try:
            workbook_order = [worksheet.title for worksheet in workbook.worksheets]
        finally:
            workbook.close()
        missing = [name for name in sheet_names if name not in workbook_order]
        if missing:
            raise ConversionError(
                "sheet_selection_invalid",
                f"worksheet selection is no longer valid: {', '.join(missing)}",
            )

        selected = set(sheet_names)
        ordered_names = [name for name in workbook_order if name in selected]
        frames = pd.read_excel(input_path, sheet_name=ordered_names, engine="openpyxl")
        html_converter = HtmlConverter()
        title = Converter._heading_text(input_path.stem)
        parts = [f"# {title}"]
        for name in ordered_names:
            frame = frames[name].rename(
                columns=lambda column: ""
                if isinstance(column, str) and re.fullmatch(r"Unnamed: \d+", column)
                else column
            )
            table = html_converter.convert_string(
                frame.to_html(index=False, na_rep="")
            ).markdown.strip()
            parts.append(f"## {Converter._heading_text(name)}\n\n{table}")
        return "\n\n".join(parts).rstrip() + "\n"

    @staticmethod
    def _heading_text(value: str) -> str:
        return re.sub(r"[\r\n]+", " ", value).strip() or "Untitled"

    @staticmethod
    def _map_conversion_error(error: Exception) -> ConversionError:
        message = str(error) or error.__class__.__name__
        lowered = message.lower()
        if "password" in lowered or "encrypted" in lowered:
            return ConversionError("password_protected", message)
        if "not a zip file" in lowered or "invalid xml" in lowered:
            return ConversionError("unreadable_file", message)
        return ConversionError("conversion_failed", message, retryable=True)
