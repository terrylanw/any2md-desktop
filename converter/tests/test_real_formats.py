from pathlib import Path
import re
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from service import ConversionError, Converter


@pytest.fixture(scope="module")
def converter():
    return Converter()


def test_converts_docx(tmp_path: Path, converter: Converter):
    source = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Quarterly Report", level=1)
    document.add_paragraph("Revenue increased.")
    document.save(source)

    result = converter.convert(str(source), "docx")

    assert "Quarterly Report" in result.markdown
    assert "Revenue increased" in result.markdown


def test_converts_xlsx(tmp_path: Path, converter: Converter):
    source = tmp_path / "budget.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Department", "Budget"])
    sheet.append(["Research", 1200])
    workbook.save(source)

    result = converter.convert(str(source), "xlsx")

    assert "Department" in result.markdown
    assert "Research" in result.markdown


def test_inspects_and_converts_selected_xlsx_sheets(
    tmp_path: Path, converter: Converter
):
    source = tmp_path / "quarterly.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Revenue", 1200])
    details = workbook.create_sheet("Details")
    details.append(["Region", "Value"])
    details.append(["North", 700])
    hidden = workbook.create_sheet("Archive")
    hidden.sheet_state = "hidden"
    hidden.append(["Secret", "Archived value"])
    workbook.create_sheet("Empty")
    workbook.save(source)
    original_bytes = source.read_bytes()

    inspection = converter.inspect_workbook(str(source), "xlsx")

    assert [sheet.name for sheet in inspection.sheets] == [
        "Summary",
        "Details",
        "Archive",
        "Empty",
    ]
    assert inspection.sheets[0].preview[1] == ["Revenue", "1200"]
    assert inspection.sheets[2].visibility == "hidden"
    assert inspection.sheets[3].is_empty is True

    result = converter.convert(str(source), "xlsx", ["Details", "Summary"])

    assert result.markdown.startswith("# quarterly\n")
    assert result.markdown.index("## Summary") < result.markdown.index("## Details")
    assert "| Metric | Value |" in result.markdown
    assert "<table" not in result.markdown
    assert "Revenue" in result.markdown
    assert "North" in result.markdown
    assert "Archived value" not in result.markdown
    assert source.read_bytes() == original_bytes


def test_xlsx_preview_is_capped(tmp_path: Path, converter: Converter):
    source = tmp_path / "large-preview.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in range(25):
        sheet.append([f"r{row}c{column}" for column in range(15)])
    sheet.cell(row=1, column=1, value="x" * 250)
    workbook.save(source)

    inspection = converter.inspect_workbook(str(source), "xlsx")
    preview = inspection.sheets[0].preview

    assert len(preview) == 20
    assert len(preview[0]) == 12
    assert len(preview[0][0]) == 200
    assert preview[0][0].endswith("...")


def test_inspection_recovers_from_incorrect_a1_dimension(
    tmp_path: Path, converter: Converter
):
    source = tmp_path / "incorrect-dimension.xlsx"
    rewritten = tmp_path / "rewritten.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Language", "Translation", "Notes"])
    sheet.append(["English", "中文", "Visible beyond A1"])
    workbook.save(source)

    with ZipFile(source, "r") as source_zip, ZipFile(rewritten, "w") as target_zip:
        for item in source_zip.infolist():
            content = source_zip.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(
                    rb'<dimension ref="[^"]+"',
                    b'<dimension ref="A1:A1"',
                    content,
                    count=1,
                )
            target_zip.writestr(item, content)

    inspection = converter.inspect_workbook(str(rewritten), "xlsx")
    inspected_sheet = inspection.sheets[0]

    assert inspected_sheet.row_count == 2
    assert inspected_sheet.column_count == 3
    assert inspected_sheet.preview[1] == ["English", "中文", "Visible beyond A1"]


def test_rejects_stale_or_empty_sheet_selection(
    tmp_path: Path, converter: Converter
):
    source = tmp_path / "selection.xlsx"
    workbook = Workbook()
    workbook.save(source)

    with pytest.raises(ConversionError) as empty:
        converter.convert(str(source), "xlsx", [])
    assert empty.value.code == "sheet_selection_invalid"

    with pytest.raises(ConversionError) as stale:
        converter.convert(str(source), "xlsx", ["Missing"])
    assert stale.value.code == "sheet_selection_invalid"


def test_selected_xlsx_does_not_emit_dataframe_placeholders(
    tmp_path: Path, converter: Converter
):
    source = tmp_path / "blank-cells.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Translations"
    sheet.append(["Name", None])
    sheet.append(["Alpha", None])
    sheet.append(["Beta", "Late value"])
    workbook.save(source)

    result = converter.convert(str(source), "xlsx", ["Translations"])

    assert "Unnamed:" not in result.markdown
    assert "NaN" not in result.markdown
    assert "Late value" in result.markdown


def test_converts_pptx(tmp_path: Path, converter: Converter):
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Launch Plan"
    slide.placeholders[1].text = "Prepare beta release"
    presentation.save(source)

    result = converter.convert(str(source), "pptx")

    assert "Launch Plan" in result.markdown
    assert "Prepare beta release" in result.markdown


def test_converts_text_pdf(tmp_path: Path, converter: Converter):
    source = tmp_path / "memo.pdf"
    document = canvas.Canvas(str(source))
    document.drawString(72, 760, "Any2MD text PDF conversion fixture")
    document.save()

    result = converter.convert(str(source), "pdf")

    assert "Any2MD text PDF conversion fixture" in result.markdown


def test_rejects_pdf_without_text(tmp_path: Path, converter: Converter):
    source = tmp_path / "blank.pdf"
    document = canvas.Canvas(str(source))
    document.showPage()
    document.save()

    with pytest.raises(ConversionError) as caught:
        converter.convert(str(source), "pdf")

    assert caught.value.code == "no_extractable_text"


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("notes.txt", "Local Markdown conversion", "Local Markdown conversion"),
        ("data.csv", "Name,Value\nAlpha,42\n", "Alpha"),
        ("page.html", "<h1>Local HTML</h1><p>Converted safely.</p>", "Local HTML"),
    ],
)
def test_converts_text_formats(
    tmp_path: Path,
    converter: Converter,
    name: str,
    content: str,
    expected: str,
):
    source = tmp_path / name
    source.write_text(content, encoding="utf-8")

    result = converter.convert(str(source))

    assert expected in result.markdown
